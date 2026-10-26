"""Export impact / upstream results to CSV, XLSX, JSON and a self-contained HTML report."""
from __future__ import annotations

import csv
import datetime as dt
import html
import io
import json
import warnings
from pathlib import Path

from . import __version__
from .graph import BREAKS, NODE_TYPE_LABELS

COLUMNS = ["impact", "confidence", "type", "object", "report", "page", "depth", "reason", "path"]
GROUP_ORDER = ["visual", "page", "report", "measure", "calc_column", "rls_role", "calc_group_item", "field_parameter",
               "pbi_model_column", "snowflake_view_column", "dbt_model_column", "snowflake_table_column", "dbt_source_column"]


def flat_rows(result: dict) -> list[dict]:
    rows = []
    groups = result.get("groups", {})
    for t in GROUP_ORDER + [k for k in groups if k not in GROUP_ORDER]:
        for e in groups.get(t, []):
            report = e.get("container", "") if t in ("visual", "page") else (e["name"] if t == "report" else "")
            page = e.get("parent", "") if t == "visual" else (e["name"] if t == "page" else "")
            rows.append({
                "impact": "breaks" if e["impact"] == BREAKS else "values may change",
                "confidence": e["confidence"], "type": NODE_TYPE_LABELS.get(t, t), "object": e["display"],
                "report": report, "page": page, "depth": e["depth"], "reason": e.get("reason", ""),
                "path": " → ".join(e.get("path", [])), "id": e["id"],
            })
    return rows


def to_csv(result: dict) -> str:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLUMNS + ["id"], extrasaction="ignore")
    w.writeheader()
    w.writerows(flat_rows(result))
    return buf.getvalue()


def to_json(result: dict) -> str:
    return json.dumps({"generator": f"LineageLens {__version__}", "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
                       **result}, indent=2, default=str)


def to_xlsx(result: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4F46E5")
    red = PatternFill("solid", fgColor="FEE2E2")
    amber = PatternFill("solid", fgColor="FEF3C7")
    ws["A1"] = "LineageLens impact report"
    ws["A1"].font = Font(bold=True, size=16)
    starts = ", ".join(s["display"] for s in result.get("start", []))
    ws.append([])
    ws.append(["Changed object", starts])
    ws.append(["Change type", result.get("change_label", "")])
    ws.append(["Generated", dt.datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append([])
    for k, v in result.get("summary", {}).items():
        ws.append([k.replace("_", " ").capitalize(), v])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60

    rows = flat_rows(result)

    def sheet(title, data, cols):
        s = wb.create_sheet(title)
        s.append([c.capitalize() for c in cols])
        for i, c in enumerate(cols, 1):
            cell = s.cell(row=1, column=i)
            cell.font, cell.fill = head, fill
        for r in data:
            s.append([r.get(c, "") for c in cols])
            f = red if r.get("impact") == "breaks" else amber
            s.cell(row=s.max_row, column=1).fill = f
        widths = {"object": 55, "path": 90, "reason": 30, "report": 22, "page": 24, "type": 22}
        for i, c in enumerate(cols, 1):
            s.column_dimensions[get_column_letter(i)].width = widths.get(c, 16)
        s.freeze_panes = "A2"
        s.auto_filter.ref = s.dimensions
        for row in s.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)

    sheet("Visuals", [r for r in rows if r["type"] in ("Visual", "Page filter", "Report filter", "Page", "Report")],
          ["impact", "confidence", "report", "page", "object", "reason", "path"])
    sheet("Measures & RLS", [r for r in rows if r["type"] not in ("Visual", "Page", "Report") and not r["id"].startswith(("dbt:", "sf:"))],
          ["impact", "confidence", "type", "object", "reason", "path"])
    sheet("Warehouse", [r for r in rows if r["id"].startswith(("dbt:", "sf:"))], ["impact", "confidence", "type", "object", "reason", "path"])
    sheet("All", rows, COLUMNS)
    out = io.BytesIO()
    with warnings.catch_warnings():
        # openpyxl opens a text-mode NamedTemporaryFile without an encoding (it only ever
        # writes ASCII XML there). Silence that third-party EncodingWarning here only.
        warnings.filterwarnings("ignore", category=EncodingWarning, module=r"openpyxl\.")
        wb.save(out)
    return out.getvalue()


CSS = """
:root{--bg:#f8fafc;--card:#fff;--fg:#0f172a;--muted:#64748b;--line:#e2e8f0;--brand:#4f46e5;--red:#dc2626;--redbg:#fef2f2;--amber:#b45309;--amberbg:#fffbeb}
@media (prefers-color-scheme:dark){:root{--bg:#0b1020;--card:#121a2e;--fg:#e2e8f0;--muted:#94a3b8;--line:#24304a;--brand:#818cf8;--red:#f87171;--redbg:#2a1215;--amber:#fbbf24;--amberbg:#2a2010}}
*{box-sizing:border-box}body{margin:0;font:14px/1.5 -apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;background:var(--bg);color:var(--fg)}
.wrap{max-width:1100px;margin:0 auto;padding:32px 24px}
header{display:flex;justify-content:space-between;align-items:flex-end;border-bottom:1px solid var(--line);padding-bottom:16px;margin-bottom:24px}
h1{margin:0;font-size:24px}h2{font-size:17px;margin:28px 0 10px}h3{font-size:14px;margin:16px 0 6px}
.logo{display:inline-flex;align-items:center;gap:8px;color:var(--brand);font-weight:700;font-size:13px;letter-spacing:.04em;text-transform:uppercase}
.muted{color:var(--muted)}code,.mono{font-family:ui-monospace,SFMono-Regular,Consolas,monospace;font-size:12.5px}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px}
.card{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:14px 16px}
.card .v{font-size:26px;font-weight:700}.card .k{color:var(--muted);font-size:12px;text-transform:uppercase;letter-spacing:.05em}
.card.red .v{color:var(--red)}
.badge{display:inline-block;border-radius:999px;padding:1px 9px;font-size:11.5px;font-weight:600;white-space:nowrap}
.b-breaks{background:var(--redbg);color:var(--red)}.b-may{background:var(--amberbg);color:var(--amber)}
.b-high{background:rgba(22,163,74,.12);color:#16a34a}.b-medium{background:rgba(234,179,8,.15);color:#a16207}.b-low{background:rgba(220,38,38,.12);color:var(--red)}
table{width:100%;border-collapse:collapse;background:var(--card);border:1px solid var(--line);border-radius:12px;overflow:hidden}
th,td{text-align:left;padding:8px 12px;border-bottom:1px solid var(--line);vertical-align:top}th{font-size:12px;color:var(--muted);text-transform:uppercase;letter-spacing:.04em}
tr:last-child td{border-bottom:0}
.report{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:14px 18px;margin-bottom:12px}
.page{border-left:3px solid var(--line);padding-left:12px;margin:10px 0}
.path{color:var(--muted);font-size:12px}
footer{margin-top:40px;color:var(--muted);font-size:12px;border-top:1px solid var(--line);padding-top:12px}
"""


def _badge(imp: str) -> str:
    return '<span class="badge b-breaks">breaks</span>' if imp == BREAKS else '<span class="badge b-may">values may change</span>'


def _conf(c: str) -> str:
    return f'<span class="badge b-{c}">{c}</span>'


def to_html(result: dict, title: str | None = None) -> str:
    e = html.escape
    s = result.get("summary", {})
    starts = ", ".join(x["display"] for x in result.get("start", []))
    title = title or f"Impact of changing {starts}"
    out = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
           f"<title>{e(title)}</title><style>{CSS}</style></head><body><div class='wrap'>"]
    out.append(f"<header><div><div class='logo'>◆ LineageLens</div><h1>{e(title)}</h1>"
               f"<div class='muted'>Change type: <b>{e(result.get('change_label', ''))}</b></div></div>"
               f"<div class='muted'>{dt.datetime.now().strftime('%d %b %Y, %H:%M')}</div></header>")
    cards = [("Impacted objects", s.get("total", 0), False), ("Breaking", s.get("breaks", 0), True),
             ("Visuals", s.get("visuals", 0), False), ("Measures", s.get("measures", 0), False),
             ("RLS roles", s.get("rls_roles", 0), s.get("rls_roles", 0) > 0), ("Reports", s.get("reports", 0), False)]
    out.append("<div class='cards'>" + "".join(f"<div class='card{' red' if red and v else ''}'><div class='k'>{e(k)}</div><div class='v'>{v}</div></div>" for k, v, red in cards) + "</div>")
    out.append("<h2>Reports, pages &amp; visuals</h2>")
    if not result.get("reports"):
        out.append("<p class='muted'>No report visuals are affected.</p>")
    for r in result.get("reports", []):
        out.append(f"<div class='report'><div style='display:flex;justify-content:space-between'><b>{e(r['name'])}</b>"
                   f"<span>{_badge(r['impact'])} <span class='muted'>{r['visual_count']} visual(s)</span></span></div>")
        if r.get("report_filter"):
            out.append(f"<div class='path'>Report-level filter affected {_badge(r['report_filter']['impact'])}</div>")
        for p in r["pages"]:
            out.append(f"<div class='page'><div><b>{e(p['name'])}</b> <span class='muted'>· {len(p['visuals'])} of {p['total_visuals']} visuals</span>"
                       + (f" · page filter {_badge(p['page_filter']['impact'])}" if p.get("page_filter") else "") + "</div>")
            for v in p["visuals"]:
                out.append(f"<div style='margin:6px 0'>{_badge(v['impact'])} {_conf(v['confidence'])} <b>{e(v['name'])}</b> "
                           f"<span class='muted'>({e(v.get('visual_type', ''))}; {e(v.get('usage', ''))})</span>"
                           f"<div class='path mono'>{e(' → '.join(v.get('path', [])))}</div></div>")
            out.append("</div>")
        out.append("</div>")
    groups = result.get("groups", {})
    for t, label in [("rls_role", "Row-level security roles"), ("measure", "Measures"), ("calc_column", "Calculated columns"),
                     ("calc_group_item", "Calculation group items"), ("field_parameter", "Field parameters"),
                     ("pbi_model_column", "Power BI model columns"), ("snowflake_view_column", "Snowflake view columns"),
                     ("dbt_model_column", "dbt model columns")]:
        lst = groups.get(t, [])
        if not lst:
            continue
        out.append(f"<h2>{e(label)} <span class='muted'>({len(lst)})</span></h2><table><tr><th>Object</th><th>Impact</th><th>Confidence</th><th>Path</th></tr>")
        for x in lst:
            out.append(f"<tr><td class='mono'>{e(x['display'])}</td><td>{_badge(x['impact'])}</td><td>{_conf(x['confidence'])}</td>"
                       f"<td class='path mono'>{e(' → '.join(x.get('path', [])))}</td></tr>")
        out.append("</table>")
    out.append(f"<footer>Generated offline by LineageLens {__version__}. Deterministic parsing of dbt, Snowflake and Power BI metadata — "
               "no AI, no data leaves your machine. Confidence: high = exact match, medium = naming-convention/Power Query rename, low = fuzzy.</footer>")
    out.append("</div></body></html>")
    return "".join(out)


def export(result: dict, fmt: str) -> tuple[bytes, str, str]:
    """Returns (content, media_type, file extension)."""
    fmt = fmt.lower()
    if fmt == "csv":
        return to_csv(result).encode("utf-8-sig"), "text/csv", "csv"
    if fmt == "json":
        return to_json(result).encode(), "application/json", "json"
    if fmt == "xlsx":
        return to_xlsx(result), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    if fmt in ("html", "htm"):
        return to_html(result).encode(), "text/html", "html"
    raise ValueError(f"unknown format {fmt}")


def write(result: dict, path: str | Path, fmt: str | None = None) -> Path:
    path = Path(path)
    fmt = fmt or path.suffix.lstrip(".") or "html"
    data, _, _ = export(result, fmt)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path
